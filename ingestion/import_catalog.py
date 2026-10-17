"""Import canonical JSON and preserve every normalized workbook table."""
import argparse
import hashlib
import json
from pathlib import Path
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parents[1]
TABLES = ('institutions', 'locations', 'sources', 'institution_sources', 'aliases', 'tags',
          'institution_tags', 'announcements', 'events', 'event_evidence', 'changes')


def platform(url, declared='website'):
    host = (urlsplit(url).hostname or '').lower().removeprefix('www.')
    if host in ('t.me', 'tg.me', 'telegram.me'): return 'telegram'
    if host in ('vk.com', 'vk.ru', 'm.vk.com'): return 'vk'
    if host in ('instagram.com', 'instagr.am'): return 'instagram'
    if any(urlsplit(url).path.rstrip('/').endswith(s) for s in ('.rss', '.atom', '/feed', '/rss', '/atom')): return 'rss'
    return declared or 'website'


def source_key(url):
    if platform(url) == 'telegram':
        parts = urlsplit(url).path.strip('/').split('/')
        if parts and parts[0] == 's': parts = parts[1:]
        return ('telegram', parts[0].casefold() if parts else '')
    return ('url', url)


def import_data(json_path, workbook=None, output=ROOT, additional_path=None):
    raw = json.loads(Path(json_path).read_text())
    if not isinstance(raw.get('institutions'), list): raise ValueError('Отсутствует таблица institutions')
    ids = [i['institution_id'] for i in raw['institutions']]
    if len(ids) != len(set(ids)): raise ValueError('Повторяются идентификаторы institutions')
    normalized = {}
    if workbook:
        from openpyxl import load_workbook
        book = load_workbook(workbook, read_only=True, data_only=True)
        for table in TABLES:
            if table not in book.sheetnames: raise ValueError(f'В XLSX отсутствует таблица {table}')
            rows = book[table].iter_rows(values_only=True)
            headers = next(rows)
            normalized[table] = [dict(zip(headers, row)) for row in rows if any(v is not None for v in row)]
        book.close()
        if set(ids) != {i['institution_id'] for i in normalized['institutions']}:
            raise ValueError('Каталоги JSON и XLSX расходятся по institution_id')
    sources = {s['source_id']: dict(s) for s in normalized.get('sources', [])}
    relations = [dict(r) for r in normalized.get('institution_sources', [])]
    if not relations:
        relations = [dict(institution_id=i['institution_id'], **{k:s.get(k) for k in ('source_id','role','officiality','identity_evidence_url')})
                     for i in raw['institutions'] for s in i.get('sources', [])]
    for i in raw['institutions']:
        for s in i.get('sources', []):
            if s['source_id'] in sources and sources[s['source_id']]['url'] != s['url']:
                raise ValueError(f"Конфликт URL: {s['source_id']}")
            sources.setdefault(s['source_id'], dict(s))
    additional = None
    if additional_path:
        additional = json.loads(Path(additional_path).read_text())
        added_ids = set()
        existing_urls = {source_key(s['url']): sid for sid, s in sources.items()}
        for s in additional.get('sources', []):
            sid = s['source_id']
            if sid in added_ids or sid in sources:
                raise ValueError(f'Повторяется дополнительный source_id: {sid}')
            if source_key(s['url']) in existing_urls:
                raise ValueError(f"Повторяется дополнительный URL: {s['url']}")
            added_ids.add(sid); sources[sid] = dict(s); existing_urls[source_key(s['url'])] = sid
        existing_relations = {(r['institution_id'], r['source_id']) for r in relations}
        for relation in additional.get('institution_sources', []):
            pair = (relation['institution_id'], relation['source_id'])
            if pair in existing_relations:
                raise ValueError(f'Повторяется дополнительная связь: {pair}')
            if not relation.get('identity_evidence_url') or not relation.get('identity_checked_on'):
                raise ValueError(f'Нет подтверждения принадлежности дополнительного источника: {pair}')
            relations.append(dict(relation)); existing_relations.add(pair)
        for sid in added_ids:
            if not any(r['source_id'] == sid for r in relations):
                raise ValueError(f'Дополнительный источник без площадки: {sid}')
    for relation in relations:
        if relation['institution_id'] not in ids or relation['source_id'] not in sources:
            raise ValueError('Нарушена связь institution_sources')
    for s in sources.values():
        if urlsplit(s['url']).scheme not in ('https','http') or not urlsplit(s['url']).hostname or urlsplit(s['url']).username or urlsplit(s['url']).password or any(ord(c)<32 for c in s['url']):
            raise ValueError(f"Неверный URL: {s['source_id']}")
        s['platform'] = platform(s['url'], s.get('platform'))
        s['institution_ids'] = list(dict.fromkeys(r['institution_id'] for r in relations if r['source_id'] == s['source_id']))
        s['roles'] = list(dict.fromkeys(r.get('role') for r in relations if r['source_id'] == s['source_id']))
    institutions = []
    for original in raw['institutions']:
        i = dict(original)
        i['source_ids'] = [r['source_id'] for r in relations if r['institution_id'] == i['institution_id']]
        i['independent'] = bool(set(i.get('tags', [])) & {'artist_run','independent'})
        i['category'] = 'artist_run' if 'artist_run' in i.get('tags',[]) else 'independent' if i['independent'] else 'institution'
        i['monitoring_priority'] = {'A':'high','B':'normal','C':'low'}.get(i.get('priority'), 'normal')
        i['districts'] = sorted({l['district'] for l in i.get('locations',[]) if l.get('district')})
        institutions.append(i)
    catalog = dict(metadata=raw['metadata'], institutions=institutions, sources=list(sources.values()),
                   institution_sources=relations, tags=normalized.get('tags', []),
                   import_provenance={'json_sha256':hashlib.sha256(Path(json_path).read_bytes()).hexdigest(),
                                      'workbook_included':bool(workbook)}, seed_events=raw.get('events',[]))
    if additional:
        catalog['import_provenance']['additional_sources_sha256'] = hashlib.sha256(Path(additional_path).read_bytes()).hexdigest()
        catalog['metadata'] = {**raw['metadata'], 'updated_on': max(raw['metadata']['updated_on'], additional['updated_on']),
            'counts': {**raw['metadata'].get('counts', {}), 'sources': len(sources), 'institution_sources': len(relations)},
            'additional_sources_count': len(additional.get('sources', []))}
    if normalized: (output/'catalog'/'normalized.json').write_text(json.dumps(normalized, ensure_ascii=False, indent=2, default=str)+'\n')
    (output/'catalog'/'catalog.json').write_text(json.dumps(catalog, ensure_ascii=False, indent=2)+'\n')
    public = {k:v for k,v in catalog.items() if k not in ('seed_events',)}
    (output/'public'/'data').mkdir(parents=True, exist_ok=True)
    (output/'public'/'data'/'catalog.json').write_text(json.dumps(public, ensure_ascii=False, indent=2)+'\n')
    print(f"Импорт: {len(institutions)} пространств, {len(sources)} источника, {len(relations)} связей. Примеры событий не опубликованы.")
    return catalog


def main():
    p = argparse.ArgumentParser(description='Импорт исходного каталога')
    p.add_argument('--json', type=Path, default=ROOT/'catalog'/'original.json')
    p.add_argument('--xlsx', type=Path, default=ROOT/'catalog'/'original.xlsx')
    p.add_argument('--additional', type=Path, default=ROOT/'catalog'/'additional-sources.json')
    args = p.parse_args()
    import_data(args.json, args.xlsx if args.xlsx.exists() else None,
                additional_path=args.additional if args.additional.exists() else None)

if __name__ == '__main__': main()
